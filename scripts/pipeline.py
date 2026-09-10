"""
FRA (Failure Root-Cause Analysis) - 전체 재현 파이프라인

이 스크립트는 원본 SFDC 클레임 데이터부터 지금까지 만든 모든 파생 데이터셋을
순서대로 재현한다. Claude Code에서 이어 작업할 때, 이 파일을 보면서
"각 산출물이 어떻게 만들어졌는지"를 한눈에 확인할 수 있다.

주의: 이 파일은 그대로 실행하는 realtime 스크립트가 아니라, 지금까지의 로직을
단계별 함수로 정리한 "재현 문서" 역할이다. 각 STEP은 독립적으로 검토·실행 가능.

이미 만들어진 중간 산출물(재실행 없이 바로 쓸 수 있음):
  - data/fra_master_v3.csv                   전체 마스터 (13,482건 x 58컬럼)
  - data/fra_front_axle_with_platform.csv    Front Axle 확정 + 카테고리 + 플랫폼그룹 (1,858건)
  - data/fra_front_axle_vendor_exact.csv     위에 공급사 매칭 추가 (exact match만, 신뢰 가능)
"""

import pandas as pd
import re
from collections import defaultdict

# ============================================================================
# STEP 1: 클레임 마스터 생성 (fra_master_v3)
# ============================================================================
# 핵심 로직: 북미 클레임은 2단계 구조(딜러→NA→Global 재클레임)로 중복 존재.
# LEFT JOIN + COALESCE로 병합, Global을 최종 판정으로 신뢰.
# 검증 기준값: 13,482건 (NA단독 4,694 + Global재클레임 8,788)

def build_claim_master(raw_csv_path: str) -> pd.DataFrame:
    raw = pd.read_csv(raw_csv_path, encoding='utf-8-sig', low_memory=False)

    case_cols = [
        'Id', 'CaseNumber', 'OriginalClaimNumber__c', 'IsFromNA__c', 'IsSentHQ__c', 'RecordType.Name',
        'FailureDate__c', 'RepairDate__c', 'ClosedDate', 'CreatedDate', 'UsageTime__c',
        'fm_Glo_Kor_ProdDate__c', 'fm_AssetRetailDate__c', 'fm_Kor_ShippedDate__c',
        'CauseCode__c', 'CauseCode2__c', 'CauseCode3__c', 'ClaimType__c', 'Status', 'Subject',
        'DescriptionOfFailure__c', 'Failure_Cause__c', 'Repair__c', 'AdminNotes__c',
        'AssetId', 'fm_ItemCode__c', 'fm_ItemName__c', 'fm_ItemNameEng__c', 'Asset.Name',
        'AccountId', 'Account.Name', 'fm_DealerShipName__c',
        'fm_TotalRequestAmount__c', 'fm_TotalApprovedAmount__c',
        'ru_TotalApprovedPartsAmount__c', 'ru_PartsTotal__c',
        'ru_TotalRequestLaborCost__c', 'ru_TotalApprovedLaborCost__c',
        'ru_TotalRequestLaborHour__c', 'ru_TotalApprovedLaborHour__c', 'ru_CountParts__c',
        'DealerComment__c', 'Description',
    ]
    claims = raw[case_cols].drop_duplicates(subset=['Id']).copy()

    na = claims[claims['RecordType.Name'] == 'North America'].copy()
    glo = claims[claims['RecordType.Name'] == 'Global'].copy()
    na['id15'] = na['Id'].astype(str).str[:15]
    glo['orig15'] = glo['OriginalClaimNumber__c'].astype(str).str[:15]

    na_ids = set(na['id15'])
    glo_from_na = glo[glo['orig15'].isin(na_ids)].copy()
    linked = set(glo_from_na['orig15'])
    na_standalone = na[~na['id15'].isin(linked)].copy()

    glo_from_na['status_rank'] = (glo_from_na['Status'] != 'Closed').astype(int)
    glo_dedup = (
        glo_from_na.sort_values(['orig15', 'status_rank', 'CreatedDate'], ascending=[True, True, False])
        .drop_duplicates(subset=['orig15'], keep='first')
    )

    rows = []
    na_std_idx = na_standalone.set_index('id15')
    for id15, r in na_std_idx.iterrows():
        d = r.to_dict()
        d['record_source'] = 'NA_standalone'
        d['na_id'] = r['Id']
        d['global_id'] = None
        d['final_id'] = r['Id']
        rows.append(d)

    for orig15, r in glo_dedup.set_index('orig15').iterrows():
        matching_na = na[na['id15'] == orig15].iloc[0]
        d = r.to_dict()
        for c in case_cols:
            if pd.isna(d.get(c)) and pd.notna(matching_na.get(c)):
                d[c] = matching_na[c]
        d['record_source'] = 'Global'
        d['na_id'] = matching_na['Id']
        d['global_id'] = r['Id']
        d['final_id'] = r['Id']
        rows.append(d)

    master = pd.DataFrame(rows)
    assert len(master) == len(na), f"검증 실패: {len(master)} != {len(na)}"
    return master


# ============================================================================
# STEP 2: 분자(고장) 정의 v4 — 4단계 재검증 끝에 확정된 최종 규칙
# ============================================================================
# 정정 이력(순서와 이유를 이해하고 사용할 것):
#   - ClaimType__c='In Stock' 필터 없음 — 판매전 결함은 강력한 제조결함 증거이며,
#     사용시간 긴데 라벨만 In Stock인 것은 데이터이관 오류로 확인(담당자 확인)
#   - ClaimType__c='Shortage' 제외 — 부품 누락, 고장 아님(실측 확인)
#   - CauseCode__c='Shipping Damage' 제외 — 진짜 운송파손. ClaimType__c='Damaged'
#     4건 중 3건은 실제 고장이었으므로 이 값으로 정확히 필터링해야 함
#   - CauseCode3__c='Product improvement campaign' 제외 — 선제적 리콜/점검
#   - CauseCode3__c='Steering' 제외 — 조향 계통(팀장 확인, 차축과 다른 시스템)
#   - CauseCode3__c='Special allowance'는 포함 유지 — 이름과 달리 실제 고장
#   - ClaimType__c='In House' 필터 없음 — 처리장소일 뿐 고장여부와 무관.
#     자사설계결함 명시 사례 포함되어 있었음

def apply_failure_definition(master: pd.DataFrame) -> pd.DataFrame:
    cond = (
        (master['ClaimType__c'] != 'Shortage')
        & (master['CauseCode__c'].notna())
        & (master['CauseCode__c'] != 'Shipping Damage')
        & (
            master['CauseCode3__c'].isna()
            | (~master['CauseCode3__c'].isin(['Product improvement campaign', 'Steering']))
        )
    )
    return master[cond]


# ============================================================================
# STEP 3: Front Axle 대상 확정
# ============================================================================
# 검증 기준값: 1,991건(Branson 제외 전) -> 1,858건(Branson 제외 후)

def get_front_axle(failure_only: pd.DataFrame) -> pd.DataFrame:
    fa = failure_only[failure_only['CauseCode__c'] == 'Front Axle'].copy()
    is_branson = fa['fm_ItemName__c'].astype(str).str.contains('Branson', case=False, na=False)
    return fa[~is_branson]


# ============================================================================
# STEP 4: 부품 카테고리 분류
# ============================================================================
# causal_parts_no를 기준 식별자로, causal_parts_name(자유텍스트, 표기 분산 심함)은
# 참고 라벨. "SUB TO <코드> <이름>" 표기는 뒤쪽 실명으로 재분류.
# 검증 기준값(1,991건): Seal 454 / Bearing 394 / Gear 198 / Washer 116 등

def _strip_sub_to(name):
    if pd.isna(name):
        return None
    n = str(name)
    if n.upper().startswith('SUB TO'):
        rest = n[6:].strip()
        tokens = rest.split(None, 1)
        return tokens[1] if len(tokens) > 1 else None
    return n


def _categorize_part(name):
    if pd.isna(name):
        return 'Unclassified(부품명 정보 없음)'
    n = str(name).upper()
    if 'CIR CLIP' in n or 'CIRCLIP' in n or 'SNAP RING' in n:
        return 'Circlip/Snap Ring'
    if 'FRONT AXLE ASS' in n or "F/A ASSEMBLY" in n:
        return 'Front Axle Assembly(전체교체)'
    if 'SEAL' in n:
        return 'Seal'
    if 'BEARING' in n:
        return 'Bearing'
    if 'GEAR' in n or 'PINION' in n:
        return 'Gear'
    if 'CYLINDER' in n:
        return 'Cylinder'
    if 'WASHER' in n:
        return 'Washer'
    if 'SPINDLE' in n:
        return 'Spindle'
    if 'MOUNTING FRAME' in n or 'SUPPORT' in n or (
        'AXLE' in n and ('ASSY' in n or 'HOUSING' in n or 'BRACKET' in n or 'FRAME' in n)
    ):
        return 'Axle Housing/Frame'
    if 'TIE ROD' in n:
        return 'Tie Rod'
    if 'JOINT' in n:
        return 'Ball Joint'
    if 'RING' in n and 'GEAR' not in n and 'BEARING' not in n:
        return 'Ring(Snap/C/O)'
    if 'BOLT' in n or 'NUT' in n:
        return 'Bolt/Nut'
    if 'CASE' in n and ('FINAL DRIVE' in n or 'GEAR' in n):
        return 'Final Drive Case'
    if 'SWITCH' in n:
        return 'Switch(Electrical)'
    if 'DIFF' in n:
        return 'Differential Assy'
    if 'SHAFT' in n:
        return 'Shaft'
    if 'CAP' in n:
        return 'Cap'
    if 'SENSOR' in n:
        return 'Sensor(Electrical)'
    if 'SPRING PIN' in n or 'PIN' in n:
        return 'Pin'
    if 'BUSH' in n:
        return 'Bush'
    if 'COVER' in n or 'SHIELD' in n:
        return 'Cover'
    if 'COUPLING' in n:
        return 'Coupling'
    if 'CHECK S/N' in n:
        return 'Admin/Note(분석제외 검토)'
    if 'HOSE' in n:
        return 'Hose'
    if 'QUICK ATTACH' in n:
        return 'Quick Attach'
    if 'COLLAR' in n or 'SLEEVE' in n:
        return 'Collar/Sleeve'
    if 'BREATHER' in n or 'DAMPER' in n or 'NIPPLE' in n:
        return 'Other Small Parts'
    if 'SHIM' in n:
        return 'Shim'
    if 'BRACKET' in n or 'STAY' in n:
        return 'Bracket/Stay'
    return 'Other/Unclassified(재검토 필요)'


def add_part_category(fa: pd.DataFrame) -> pd.DataFrame:
    fa = fa.copy()
    fa['name_resolved'] = fa['causal_parts_name'].apply(_strip_sub_to)
    fa['part_category'] = fa['name_resolved'].apply(_categorize_part)
    fa.loc[fa['causal_parts_no'].isna(), 'part_category'] = 'No Part Info(부품정보 없음)'
    return fa


# ============================================================================
# STEP 5: 플랫폼 그룹 매핑 (웹사이트 스펙 검토 파일 기준, "Group" 공식 정보 사용)
# ============================================================================
# 주의: 2025_통합기종명.xlsx는 불완전(T264/T754/T654 등 다수 기종 누락 확인됨).
# 웹사이트_스펙_검토_260629.xlsx의 "트랙터(북미)" 시트, Group 행이 훨씬 완전함.
# 검증 기준값: 89.4% 매칭 (1,858건 중 1,661건)

def build_platform_mapping(spec_xlsx_path: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(spec_xlsx_path, data_only=True)
    ws = wb['트랙터(북미)']
    rows = list(ws.iter_rows(values_only=True))
    model_row = [r for r in rows if r[0] == 'Model Number'][0]
    group_row = [r for r in rows if r[0] == 'Group'][0]

    model_to_group = {}
    for i, v in enumerate(model_row):
        if v and str(v).strip() and i > 0:
            model = str(v).replace(' - 단종', '').strip().upper()
            grp = group_row[i] if i < len(group_row) else None
            model_to_group[model] = grp if grp and grp not in ('-', 'None') else f"Solo_{model}"
    return model_to_group


def add_platform_group(fa: pd.DataFrame, model_to_group: dict) -> pd.DataFrame:
    fa = fa.copy()

    def extract_model(code, name):
        for src in [name, code]:
            if pd.isna(src):
                continue
            s = str(src).upper()
            m = re.match(r'^(T\d{2,4}C?|\d{3,4})', s)
            if m:
                cand = m.group(1)
                if cand in model_to_group:
                    return cand
                base = cand.rstrip('C')
                if base in model_to_group:
                    return base
        return None

    fa['model_matched'] = fa.apply(lambda r: extract_model(r['fm_ItemCode__c'], r['fm_ItemName__c']), axis=1)
    fa['platform_group'] = fa['model_matched'].map(model_to_group)
    return fa


# ============================================================================
# STEP 6: 부품-공급사 매칭 (exact match만 사용, 신뢰 가능)
# ============================================================================
# 중요: prefix 매칭(앞 N자리만 비교)은 사용하지 않는다. 부품코드 마지막
# 자리(예: ...80 -> ...81 -> ...82)와 접미 알파벳(A, B, C...)은 리비전이 아니라
# "호환 안 되면 파생되는 별개 부품"이다(구매팀 확인). 80과 81은 실제로 다른
# 부품(스펙이 다를 수 있음)이며, 하나로 묶으면 서로 무관한 부품의 공급사를
# 뒤섞는 오류가 생긴다(과거 시도에서 97.8% 매칭까지 나왔으나 착시였음).
# 정확 매칭만 사용 시 매칭률 67.6%가 정직한 수치.
#
# 참고: 리비전 계보(80->81->82) 추적은 "설계변경 효과 검증(FRACAS Verify)"에는
# 유효하나, "공급사가 누구인가"라는 목적에는 부적합.

def build_vendor_mapping(vendor_xlsx_path: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(vendor_xlsx_path, data_only=True)
    ws = wb['최종본']
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    exact_map = defaultdict(set)
    for r in rows:
        code = str(r[2]).strip()
        if code:
            exact_map[code].add(r[5])
    return dict(exact_map)


def add_vendor_match(fa: pd.DataFrame, vendor_map: dict) -> pd.DataFrame:
    fa = fa.copy()
    fa['vendors'] = fa['causal_parts_no'].astype(str).str.strip().map(lambda x: vendor_map.get(x))
    fa['n_vendors'] = fa['vendors'].apply(lambda x: len(x) if x else 0)
    return fa


# ============================================================================
# 알려진 미해결 이슈 (파킹 중, Claude Code에서 이어갈 때 참고)
# ============================================================================
# 1. 부품-공급사 매핑표(옥천/익산 기준)에 없는 부품 87개(427건, 32.4%)
#    - 시기(생산연도/접수연도)와는 무관함(확인됨)
#    - 특정 플랫폼(Group/Model 4=T264, Group/Model 10=T494/T574)에 집중
#    - 상위 6개 부품(와셔 16709510081 등)이 미매칭의 64% 차지
#    - 원인 불명: 북미 전용 조달이라 매핑표 대상이 아닌지, 매핑표 자체가 누락된
#      것인지 미확인. 사용자가 추가 확인 중.
# 2. 부품번호 리비전 계보(예: 16709510080 -> 81 -> 82) 추적
#    - "설계변경이 실제로 효과가 있었는가"를 보는 축. 아직 미착수.
# 3. 협력사 로트 추적 (특정 클레임이 어느 협력사 납품분인지)
#    - ClaimId의 사이트(1100=옥천/2100=익산) 확인 후 양산구매팀 별도 데이터 필요
